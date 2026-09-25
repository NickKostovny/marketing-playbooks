#!/usr/bin/env python3
"""All deliverables -> one .xlsx (Google Sheets imports it with tabs intact).
Tabs: Read me | Action list | Accounts | HOLD accounts | Unify spec | People (all, gated) | Ads
Run after build_handoff.py.  Output: linkedin-ad-engagers-<date>.xlsx
"""
import csv, json, pathlib, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).parent
DATE = datetime.date.today().isoformat()
OUT = HERE / f"linkedin-ad-engagers-{DATE}.xlsx"

def rows(fn):
    with open(HERE / fn, newline="", encoding="utf-8") as f: return list(csv.reader(f))

def add_sheet(wb, title, data, widths=None, wrap_cols=()):
    ws = wb.create_sheet(title)
    for r in data: ws.append(r)
    if data:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(data[0], 1):
            w = (widths or {}).get(h)
            if w is None:
                longest = max((len(str(r[i-1])) for r in data[1:51] if i-1 < len(r)), default=8)
                w = min(max(10, longest + 2), 60)
            ws.column_dimensions[get_column_letter(i)].width = w
            if h in wrap_cols:
                for cell in ws[get_column_letter(i)][1:]: cell.alignment = Alignment(wrap_text=True, vertical="top")
    return ws

ads = json.loads((HERE / "ads.json").read_text())
gated = rows("people-gated.csv"); contacts = rows("contacts-P0-P2.csv"); accounts = rows("accounts.csv")
spec = rows("unify-accounts-spec.csv"); hold = rows("unify-accounts-HOLD-do-not-run.csv")
done = [a for a in ads if a["scraped"]]; pend = [a for a in ads if not a["scraped"]]

wb = Workbook(); ws = wb.active; ws.title = "Read me"
notes = [
 ["LinkedIn ad engagers -> SDR handoff", f"Built {DATE}"],
 [],
 ["What this is", "Everyone who reacted to or commented on our historical LinkedIn ads, deduped across ads, gated by deal status and the persona SOP, with Clay emails only where a row is actionable."],
 ["Ads covered", f"{len(done)} of {len(ads)} ads scraped; {sum(a['reactions_captured'] or 0 for a in done)} reactions; 0 comments on any ad."
   + (f" Pending: {', '.join(a['id'] for a in pend)}." if pend else "")],
 ["Read this first", "Read the priorities before sending anything. Engagement lists are account signals, not a contact list; most engagers will not clear a persona SOP."],
 [],
 ["Tabs", ""],
 ["Action list", "The short list: everyone in priority P0, P1 or P2. P0 = their company has an ACTIVE deal -> route to the deal owner, never cold outbound. P1 = clears the persona SOP (nobody did). P2 = worth a look (PD/MSAT-type role below the seniority bar, engaged more than one ad, commented, or a closed-lost account to revive). The priority column says which. work_email is filled only where Clay verified it; email_status says why otherwise."],
 ["Accounts", "One row per employer parsed from headlines, with deal status and engager names. employer_confidence = 'low' means the employer was guessed from a trailing headline segment."],
 ["HOLD accounts", "Accounts with ACTIVE HubSpot deals. Do not run these in outbound. Send to the deal owner as 'someone at your account reacted to our ads'."],
 ["Unify spec", "Account-level rows in our outbound tool's account template (20 columns). Persona SOP titles are verbatim on every row so the outbound tool surfaces real buyers, not the engagers."],
 ["People (all, gated)", "Every unique person with: ads engaged, reactions, SOP verdict (rule-based on headline = hypothesis), deal status, priority, enrich_flag."],
 ["Ads", "The ads themselves: short link, post URL, publish date (decoded from the post id), reactions shown vs captured."],
 [],
 ["Column glossary", ""],
 ["priority", "P0 active-deal account | P1 clears SOP | P2 worth a look | P4 no action"],
 ["sop_verdict", "PASS | FAIL-dept | FAIL-seniority | FAIL-both, derived from the LinkedIn headline only"],
 ["inferred_*", "Parsed from the headline. Hypotheses, not verified data. verified fields come from Clay."],
 ["strong_reaction", "Any reaction other than a plain like (celebrate, love, support, insightful). Reaction type is the intent tell."],
 ["deal_status", "From the CRM on the day of the build. UNCHECKED = account not looked up."],
]
for r in notes: ws.append(r)
ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 120
for row in ws.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    if row[0].value and not row[1].value: row[0].font = Font(bold=True, size=12)
ws["A1"].font = Font(bold=True, size=14)

add_sheet(wb, "Action list", contacts, wrap_cols=("sop_flag", "email_status", "deal_status", "job_title"))
add_sheet(wb, "Accounts", accounts, wrap_cols=("engagers", "deal_status", "routing"))
add_sheet(wb, "HOLD accounts", hold, wrap_cols=("priority_reason", "evidence_snippet", "notes_for_unify", "deal_status"))
add_sheet(wb, "Unify spec", spec, wrap_cols=("priority_reason", "evidence_snippet", "target_titles_include", "target_titles_exclude", "linkedin_search_hint", "notes_for_unify"))
add_sheet(wb, "People (all, gated)", gated, wrap_cols=("headline", "sop_reason", "deal_status", "routing", "ad_topics"))
adcols = ["short", "post_url", "urn_type", "id", "published", "topic", "scraped", "reactions_headline", "reactions_captured", "comments_headline", "comments_captured"]
add_sheet(wb, "Ads", [adcols] + [[a.get(c, "") for c in adcols] for a in ads])
wb.save(OUT); print(f"wrote {OUT.name}: {len(contacts)-1} contacts, {len(accounts)-1} accounts, {len(hold)-1} hold, {len(spec)-1} spec, {len(gated)-1} people, {len(ads)} ads")
